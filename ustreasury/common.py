"""Shared context + drawing/explanation/chart helpers for the Treasury Analyzer."""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.chart import BarChart, LineChart
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from bond.config import Palette
from dcf.utils import Refs, col_letter  # noqa: F401
from .config import APP_NAME, APP_VERSION, BAD_WORDS, GOOD_WORDS, WARN_WORDS


@dataclass
class Context:
    data: dict
    refs: Refs
    sheets: dict = field(default_factory=dict)

    def sheet(self, name):
        return self.sheets[name]


def title_block(sh, title, subtitle, *, last_col, row=1):
    sh.hide_gridlines()
    sh.merge(row, 1, row, last_col)
    sh.put(row, 1, title, role="title")
    sh.row_height(row, 28)
    sh.merge(row + 1, 1, row + 1, last_col)
    sh.put(row + 1, 1, subtitle, role="subtitle")
    sh.merge(row + 2, 1, row + 2, last_col)
    sh.put(row + 2, 1, f"{APP_NAME}  -  v{APP_VERSION}", role="note")
    return row + 3


def section(sh, row, text, *, c1=1, c2=1, role="section"):
    if c2 > c1:
        sh.merge(row, c1, row, c2)
    sh.put(row, c1, text, role=role)
    sh.row_height(row, 18)
    return row + 1


def nav_bar(sh, row, *, exclude=None, start_col=1, per_row=6):
    from .config import SHEET_ORDER
    exclude = exclude or set()
    targets = [s for s in SHEET_ORDER if s not in exclude]
    col = start_col
    for i, name in enumerate(targets):
        if i and i % per_row == 0:
            row += 1
            col = start_col
        c = sh.put(row, col, name, role="section", align="c")
        c.hyperlink = f"#'{name}'!A1"
        col += 1
    return row + 1


def explain_box(sh, row, lines, *, last_col, title="How to read this sheet"):
    sh.merge(row, 1, row, last_col)
    sh.put(row, 1, title, role="explain", bold=True)
    sh.row_height(row, 16)
    r = row + 1
    for ln in lines:
        sh.merge(r, 1, r, last_col)
        sh.put(r, 1, ln, role="explain")
        sh.row_height(r, max(14, 13 + 12 * (len(ln) // 120)))
        r += 1
    return r + 1


def legend(sh, row, last_col):
    sh.put(row, 1, "Legend:", role="label_b")
    sh.put(row, 2, "Input (type here)", role="input_c")
    sh.put(row, 4, "Formula", role="formula_l")
    sh.put(row, 5, "Key output", role="output_l")
    if last_col >= 7:
        sh.put(row, 7, "blue = input - black = formula - green = key output", role="note_l")
        sh.merge(row, 7, row, last_col)
    return row + 1


def define_name(sh, name, row, col):
    wb = sh.ws.parent
    coord = f"'{sh.name}'!${col_letter(col)}${row}"
    wb.defined_names[name] = DefinedName(name, attr_text=coord)
    return name


def cell_comment(sh, row, col, text):
    c = sh.ws.cell(row=row, column=col)
    cm = Comment(text, "Treasury Analyzer")
    cm.width, cm.height = 280, 120
    c.comment = cm


def dropdown(sh, row, col, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"',
                        allow_blank=True, showErrorMessage=True)
    sh.ws.add_data_validation(dv)
    dv.add(sh.ws.cell(row=row, column=col))


def traffic_light(sh, cell_range, anchor):
    def rule(words):
        return ["OR(" + ",".join(f'ISNUMBER(SEARCH("{w}",{anchor}))' for w in words) + ")"]
    for words, fc, fl in ((GOOD_WORDS, Palette.GOOD, Palette.GOOD_FILL),
                          (WARN_WORDS, Palette.WARN, Palette.WARN_FILL),
                          (BAD_WORDS, Palette.BAD, Palette.BAD_FILL)):
        sh.ws.conditional_formatting.add(
            cell_range, FormulaRule(formula=rule(words), stopIfTrue=True,
                                    font=Font(name="Calibri", size=10, bold=True, color=fc),
                                    fill=PatternFill(start_color=fl, end_color=fl, fill_type="solid")))


def bar_chart(sh, anchor, title, *, cats_ref, data_ref, width=15, height=8):
    ch = BarChart()
    ch.type = "col"
    ch.title = title
    ch.height = height
    ch.width = width
    ch.legend = None
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    sh.ws.add_chart(ch, anchor)


def line_chart(sh, anchor, title, *, cats_ref, data_ref, width=15, height=8,
               x_title="Maturity", y_title="Yield"):
    ch = LineChart()
    ch.title = title
    ch.height = height
    ch.width = width
    ch.legend = None
    ch.x_axis.title = x_title
    ch.y_axis.title = y_title
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    ch.marker = True
    sh.ws.add_chart(ch, anchor)
