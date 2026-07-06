"""Shared context + drawing / naming / conditional-format helpers for the macro system."""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from dcf.utils import Refs, col_letter  # noqa: F401
from .config import APP_NAME, APP_VERSION, SHEET_ORDER

# semantic signal colours (font, fill)
GREEN = ("1F7A3D", "C6EFCE")
LGREEN = ("1F7A3D", "E3F2E1")
AMBER = ("9C6500", "FFEB9C")
RED = ("9C0006", "FFC7CE")
GREY = ("595959", "EDEDED")
BLUE = ("2E5984", "DDEBF7")

# conditional-format word maps (exact-match cell equals word -> colour)
CF_FED = [("Hawkish", RED), ("Mildly hawkish", AMBER), ("Neutral", GREY),
          ("Mildly dovish", LGREEN), ("Dovish", GREEN)]
CF_TONE = [("Risk-on", GREEN), ("Neutral", GREY), ("Risk-off", RED)]
CF_INFL = [("Inflationary", RED), ("Disinflationary", GREEN), ("Neutral", GREY)]
CF_GROW = [("Growth+", GREEN), ("Growth-", RED), ("Neutral", GREY)]
CF_HOT = [("Hot", AMBER), ("Cold", BLUE), ("Neutral", GREY)]
CF_CONF = [("High", GREEN), ("Medium", AMBER), ("Low", GREY)]


@dataclass
class Context:
    data: dict
    refs: Refs
    sheets: dict = field(default_factory=dict)

    def sheet(self, name):
        return self.sheets[name]


def title_block(sh, title, subtitle, *, last_col, row=1):
    sh.hide_gridlines()
    sh.merge(row, 1, row, last_col)
    sh.put(row, 1, title, role="title")
    sh.row_height(row, 26)
    sh.merge(row + 1, 1, row + 1, last_col)
    sh.put(row + 1, 1, subtitle, role="subtitle")
    sh.merge(row + 2, 1, row + 2, last_col)
    sh.put(row + 2, 1, f"{APP_NAME}  -  v{APP_VERSION}", role="note")
    return row + 3


def section(sh, row, text, *, c1=1, c2=1, role="section"):
    if c2 > c1:
        sh.merge(row, c1, row, c2)
    sh.put(row, c1, text, role=role)
    sh.row_height(row, 18)
    return row + 1


def nav_bar(sh, row, *, exclude=None, start_col=1, per_row=6):
    exclude = exclude or set()
    targets = [s for s in SHEET_ORDER if s not in exclude]
    col = start_col
    for i, name in enumerate(targets):
        if i and i % per_row == 0:
            row += 1
            col = start_col
        c = sh.put(row, col, name, role="section", align="c")
        c.hyperlink = f"#'{name}'!A1"
        col += 1
    return row + 1


def _box(sh, row, lines, *, last_col, title):
    if isinstance(lines, str):
        lines = [lines]
    sh.merge(row, 1, row, last_col)
    sh.put(row, 1, title, role="explain", bold=True)
    sh.row_height(row, 16)
    r = row + 1
    for ln in lines:
        sh.merge(r, 1, r, last_col)
        sh.put(r, 1, ln, role="explain")
        sh.row_height(r, max(14, 13 + 12 * (len(ln) // 118)))
        r += 1
    return r + 1


def explain_box(sh, row, lines, *, last_col, title="What this sheet does"):
    return _box(sh, row, lines, last_col=last_col, title=title)


def interp(sh, row, lines, *, last_col, title="How to read it"):
    return _box(sh, row, lines, last_col=last_col, title=title)


def legend(sh, row, last_col):
    sh.put(row, 1, "Legend:", role="label_b")
    sh.put(row, 2, "Input (type here)", role="input_c")
    sh.put(row, 4, "Formula", role="formula_l")
    sh.put(row, 5, "Key output", role="output_l")
    if last_col >= 7:
        sh.put(row, 7, "yellow = input · black = formula · green = key output", role="note_l")
        sh.merge(row, 7, row, last_col)
    return row + 1


def define_name(sh, name, r1, c1, r2=None, c2=None):
    """Register a workbook named range: a single cell, or a rectangle if r2/c2 given."""
    wb = sh.ws.parent
    if r2 is None:
        coord = f"'{sh.name}'!${col_letter(c1)}${r1}"
    else:
        coord = f"'{sh.name}'!${col_letter(c1)}${r1}:${col_letter(c2)}${r2}"
    wb.defined_names[name] = DefinedName(name, attr_text=coord)
    return name


def cell_comment(sh, row, col, text):
    c = sh.ws.cell(row=row, column=col)
    cm = Comment(text, "Macro System")
    cm.width, cm.height = 300, 130
    c.comment = cm


def dropdown(sh, row, col, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"',
                        allow_blank=True, showErrorMessage=True)
    sh.ws.add_data_validation(dv)
    dv.add(sh.ws.cell(row=row, column=col))


def dropdown_ref(sh, row, col, name):
    """Dropdown whose source is a named range (so the list stays in one place)."""
    dv = DataValidation(type="list", formula1=f"={name}", allow_blank=True, showErrorMessage=True)
    sh.ws.add_data_validation(dv)
    dv.add(sh.ws.cell(row=row, column=col))


def word_cf(sh, cell_range, pairs):
    """Colour a range by exact-match text: pairs = [(word, (font_hex, fill_hex)), ...]."""
    for word, (fc, fl) in pairs:
        sh.ws.conditional_formatting.add(cell_range, CellIsRule(
            operator="equal", formula=[f'"{word}"'],
            font=Font(name="Calibri", size=10, bold=True, color=fc),
            fill=PatternFill(start_color=fl, end_color=fl, fill_type="solid")))


def arrow(x):
    """Excel expression fragment mapping a numeric expression `x` to an arrow glyph."""
    return (f'IF({x}>=1.5,"⇈",IF({x}>=0.5,"↑",IF({x}>-0.5,"↔",IF({x}>-1.5,"↓","⇊"))))')
