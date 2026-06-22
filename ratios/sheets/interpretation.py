"""Ratio Interpretation Guide (Sheet 14): a static reference for every ratio."""

from __future__ import annotations

from .. import common
from ..config import GUIDE

COLS = ["Ratio", "Formula", "What it measures", "Strong", "Average", "Weak",
        "Investment meaning"]
WIDTHS = [22, 30, 34, 12, 12, 12, 50]


def build(sh, ctx):
    last = len(COLS)
    common.title_block(sh, "RATIO INTERPRETATION GUIDE",
                       "Definition, formula, ranges and investment meaning for every ratio",
                       last_col=last)
    common.nav_bar(sh, 5, exclude={"Home"})
    r = 7
    for j, (name, w) in enumerate(zip(COLS, WIDTHS), start=1):
        sh.put(r, j, name, role="colhdr" if j != 1 else "colhdr")
        sh.col_width(j, w)
    sh.freeze(f"A{r+1}")
    r += 1
    for row in GUIDE:
        name, formula, measures, strong, avg, weak, meaning = row
        sh.put(r, 1, name, role="label_b")
        sh.put(r, 2, formula, role="formula_l")
        sh.put(r, 3, measures, role="note_l")
        sh.put(r, 4, strong, role="status", align="c")
        sh.put(r, 5, avg, role="status", align="c")
        sh.put(r, 6, weak, role="status", align="c")
        sh.put(r, 7, meaning, role="note", align="lw")
        sh.row_height(r, 30)
        r += 1
    # colour the strong/average/weak columns as a static legend (green/amber/red)
    from openpyxl.styles import PatternFill, Font
    from ..config import Palette
    for col, fill, font in ((4, Palette.GOOD_FILL, Palette.GOOD),
                            (5, Palette.WARN_FILL, Palette.WARN),
                            (6, Palette.BAD_FILL, Palette.BAD)):
        for rr in range(8, r):
            c = sh.ws.cell(row=rr, column=col)
            c.fill = PatternFill(start_color=fill, end_color=fill, fill_type="solid")
            c.font = Font(name="Calibri", size=9, color=font, bold=True)
    return sh
