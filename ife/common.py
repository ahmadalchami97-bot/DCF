"""Shared context + drawing/visual helpers for the forecasting engine.

The model's timeline is dynamic: the workbook auto-detects the last reported year
(by counting the filled revenue cells) and begins forecasting the year after it.
``FirstYear`` and ``LastActual`` are workbook named ranges so every label and
formula can key off them transparently.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from dcf.utils import Refs, col_letter  # noqa: F401  (Refs used as a type)
from .config import (ACTUAL_YEARS, APP_NAME, APP_VERSION, BAD_WORDS, FIRST_YEAR,
                     GOOD_WORDS, HIST_SLOTS, HORIZON, Palette, WARN_WORDS)

LABEL_COL = 1
FIRST_COL = 2          # first year column on every sheet
FY_FMT = '"FY"0'       # show the number 2025 as "FY2025"


@dataclass
class Context:
    data: dict
    refs: Refs
    sheets: dict = field(default_factory=dict)

    @property
    def first_year(self):
        return self.data.get("first_year", FIRST_YEAR)

    @property
    def actual_years(self):
        return self.data.get("actual_years", ACTUAL_YEARS)

    @property
    def horizon(self):
        return self.data.get("horizon", HORIZON)

    @property
    def hist_slots(self):
        return HIST_SLOTS

    @property
    def last_actual_year(self):
        """Python-side last actual year (for static notes only; cells use LastActual)."""
        return self.first_year + self.actual_years - 1

    def sheet(self, name):
        return self.sheets[name]


# --- drawing ----------------------------------------------------------------
def title_block(sh, title, subtitle, *, last_col, row=1):
    sh.hide_gridlines()
    sh.merge(row, 1, row, last_col)
    sh.put(row, 1, title, role="title")
    sh.row_height(row, 28)
    sh.merge(row + 1, 1, row + 1, last_col)
    sh.put(row + 1, 1, subtitle, role="subtitle")
    sh.merge(row + 2, 1, row + 2, last_col)
    sh.put(row + 2, 1, f"{APP_NAME}  -  v{APP_VERSION}", role="note")
    return row + 3


def section(sh, row, text, *, c1=1, c2=1, role="header"):
    if c2 > c1:
        sh.merge(row, c1, row, c2)
    sh.put(row, c1, text, role=role)
    sh.row_height(row, 18 if role != "header" else 20)
    return row + 1


def nav_bar(sh, row, *, exclude=None, start_col=1, per_row=10):
    from .config import SHEET_ORDER
    exclude = exclude or set()
    targets = [s for s in SHEET_ORDER if s not in exclude]
    col = start_col
    for i, name in enumerate(targets):
        if i and i % per_row == 0:
            row += 1
            col = start_col
        c = sh.put(row, col, name, role="subheader", align="c")
        c.hyperlink = f"#'{name}'!A1"
        sh.col_width(col, 13)
        col += 1
    return row + 1


def units_note(ctx):
    m = ctx.data["meta"]
    return (f"All figures in {m.get('currency','USD')} {m.get('units','millions')} unless "
            "stated.  Colour code:  blue = input  -  black = formula  -  green = key output  -  "
            "red = warning.")


def legend(sh, row, last_col):
    """A compact Actual-vs-Forecast / colour legend so the sheet is self-explaining."""
    sh.put(row, 1, "Legend:", role="label_b")
    sh.put(row, 2, "Actual (reported input)", role="actual_tag")
    sh.put(row, 4, "Forecast (model output)", role="fcst_tag")
    sh.put(row, 6, "blue = input  -  black = formula  -  green = output", role="note_l")
    if last_col > 8:
        sh.merge(row, 6, row, last_col)
    return row + 1


# --- timeline ---------------------------------------------------------------
def define_name(sh, name, row, col):
    """Create a workbook-level named range pointing at one cell; return its ref."""
    wb = sh.ws.parent
    coord = f"'{sh.name}'!${col_letter(col)}${row}"
    wb.defined_names[name] = DefinedName(name, attr_text=coord)
    return name


def shade_by_status(sh, cell_range, year_row_anchor):
    """Conditionally shade a block: actual columns vs empty 'spare' slots.

    *year_row_anchor* is a column-relative, row-absolute ref (e.g. 'B$9') to the
    fiscal-year number row, compared against the LastActual named range.
    """
    actual = PatternFill(start_color=Palette.ACTUAL_FILL, end_color=Palette.ACTUAL_FILL, fill_type="solid")
    spare = PatternFill(start_color=Palette.SPARE_FILL, end_color=Palette.SPARE_FILL, fill_type="solid")
    sh.ws.conditional_formatting.add(
        cell_range, FormulaRule(formula=[f"{year_row_anchor}>LastActual"], stopIfTrue=True, fill=spare))
    sh.ws.conditional_formatting.add(
        cell_range, FormulaRule(formula=[f"{year_row_anchor}<=LastActual"], stopIfTrue=False, fill=actual))


# --- dropdowns --------------------------------------------------------------
def dropdown(sh, row, col, options):
    formula = '"' + ",".join(options) + '"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=False, showErrorMessage=True)
    sh.ws.add_data_validation(dv)
    dv.add(sh.ws.cell(row=row, column=col))


# --- status colouring -------------------------------------------------------
def traffic_light(sh, cell_range, anchor):
    def rule(words):
        return ["OR(" + ",".join(f'ISNUMBER(SEARCH("{w}",{anchor}))' for w in words) + ")"]
    for words, fc, fl in ((GOOD_WORDS, Palette.GOOD, Palette.GOOD_FILL),
                          (WARN_WORDS, Palette.WARN, Palette.WARN_FILL),
                          (BAD_WORDS, Palette.BAD, Palette.BAD_FILL)):
        sh.ws.conditional_formatting.add(
            cell_range, FormulaRule(formula=rule(words), stopIfTrue=True,
                                    font=Font(name="Calibri", size=10, bold=True, color=fc),
                                    fill=PatternFill(start_color=fl, end_color=fl, fill_type="solid")))


# --- charts -----------------------------------------------------------------
def line_chart(sh, anchor, title, *, rows, cat_row, first_col, last_col, width=22, height=8):
    ws = sh.ws
    ch = LineChart()
    ch.title = title
    ch.height = height
    ch.width = width
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    data = Reference(ws, min_col=LABEL_COL, max_col=last_col, min_row=rows[0], max_row=rows[-1])
    ch.add_data(data, from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=first_col, max_col=last_col, min_row=cat_row, max_row=cat_row))
    ws.add_chart(ch, anchor)
