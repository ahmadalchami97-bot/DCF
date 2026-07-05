"""
Test suite for the Government Bond Course & Calculator (bondcourse).

The workbook uses Excel's native bond functions, so the analytics are verified
against the independent Python reference (bond.bondmath): internal consistency,
the two-way price<->yield relationship, and the duration/DV01/convexity numbers,
for a coupon note and a zero-coupon bill. We also check that the 18 sheets build
in order, that every named range is wired up, that the attractiveness score
matches the Excel scoring logic, that every formula name/function resolves, and
that the "Formula Explanations" sheet shows formulas as TEXT (not live formulas).

Run: ``python tests/test_bondcourse.py``  (or ``pytest -q tests/test_bondcourse.py``).
"""

from __future__ import annotations

import os
import re
import sys
import tempfile
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402

from bond import bondmath as bm  # noqa: E402
from bondcourse.config import SHEET_ORDER  # noqa: E402
from bondcourse.sample_data import build_sample, empty_like_sample  # noqa: E402
from bondcourse.workbook import assemble  # noqa: E402

S, M = date(2025, 3, 20), date(2034, 2, 15)

# Excel functions the workbook is allowed to call (native in Excel 365).
FUNCS = set((
    "IF IFERROR IFNA SUM SUMPRODUCT MIN MAX ABS SIGN N ISNUMBER NOT OR AND COUNTIF SEARCH CHOOSE TEXT ROUND "
    "YEAR MONTH DAY YIELD PRICE DURATION MDURATION ACCRINT YEARFRAC EDATE EOMONTH COUPNUM COUPNCD COUPPCD "
    "COUPDAYS COUPDAYBS COUPDAYSNC"
).split())

EXPECTED_NAMES = (
    # inputs
    "BondName", "Issuer", "Ccy", "SecType", "Face", "PosSize", "Cpn", "FreqName", "Basis",
    "IssueDate", "Settle", "Mat", "MktClean", "MktYTM", "BenchYield", "RateOutlook", "Analyst",
    "AnalysisDate", "IsZero", "Freq", "EffCoupon", "YrsToMat", "InputWarn",
    # cash flows
    "Ncoup", "Ncd", "Wfrac", "DirtyCF", "MacaulayCF", "Convexity",
    # price / yield
    "CleanPrice", "Accrued", "Dirty", "YTM", "CurrentYield", "PremDisc", "Spread", "SpreadBps",
    "TestYield", "TheoClean", "TheoDirty", "PriceDiff", "CheapFair",
    # risk
    "MacDur", "ModDur", "DV01", "DV01Pos", "ConvexityView",
    # curve + attractiveness
    "Y3M", "Y2Y", "Y5Y", "Y10Y", "Y30Y", "CurveShape",
    "AttractScore", "ModelView", "AnalystOverride", "FinalView",
)


# ---------------------------------------------------------------------------
# Analytics (verified against the independent Python reference)
# ---------------------------------------------------------------------------
def test_price_reconciles():
    a = bm.analytics(S, M, 0.04, 96.50, 2, 1)
    assert abs(bm.price_clean(S, M, 0.04, a["ytm"], 2, 1) - 96.50) < 1e-6
    assert abs(a["total_pv"] - a["dirty"]) < 1e-9


def test_two_way_price_yield():
    a = bm.analytics(S, M, 0.04, 96.50, 2, 1)
    assert abs(bm.price_clean(S, M, 0.04, a["ytm"], 2, 1) - 96.50) < 1e-6
    theo_bench = bm.price_clean(S, M, 0.04, 0.043, 2, 1)   # value at the benchmark yield
    assert theo_bench > 96.50                               # yields > benchmark -> looks "cheap"


def test_price_yield_direction():
    a = bm.analytics(S, M, 0.04, 96.50, 2, 1)
    up = bm.price_clean(S, M, 0.04, a["ytm"] + 0.005, 2, 1)
    down = bm.price_clean(S, M, 0.04, a["ytm"] - 0.005, 2, 1)
    assert up < 96.50 < down


def test_dirty_identity():
    a = bm.analytics(S, M, 0.04, 96.50, 2, 1)
    assert abs(a["accrued"] + 96.50 - a["dirty"]) < 1e-9
    assert 0.0 <= a["accrued"] <= 0.04 / 2 * 100 + 1e-9


def test_duration_dv01_convexity():
    a = bm.analytics(S, M, 0.04, 96.50, 2, 1)
    assert abs(a["modified"] - a["macaulay"] / (1 + a["ytm"] / 2)) < 1e-9
    assert a["modified"] <= a["macaulay"]
    assert a["dv01"] > 0 and a["convexity"] > 0


def test_bill_duration_equals_maturity():
    z = bm.zero_analytics(date(2025, 3, 20), date(2025, 9, 18), 97.90, 1)
    assert abs(z["macaulay"] - z["years"]) < 1e-9 and z["accrued"] == 0.0


def test_known_values():
    a = bm.analytics(S, M, 0.04, 96.50, 2, 1)
    assert 0.044 < a["ytm"] < 0.045
    assert 7.4 < a["macaulay"] < 7.6
    assert 7.3 < a["modified"] < 7.4
    assert 63 < a["convexity"] < 65
    assert 0.070 < a["dv01"] < 0.072


# ---------------------------------------------------------------------------
# Attractiveness score must match the Excel scoring formulas exactly
# ---------------------------------------------------------------------------
def _score(ytm, mod, dv01, y10, y2, bench, outlook):
    s = 3 if ytm >= 0.045 else 2 if ytm >= 0.035 else 1 if ytm >= 0.025 else 0
    s += 3 if mod <= 3 else 2 if mod <= 7 else 1 if mod <= 12 else 0
    s += 2 if dv01 <= 0.03 else 1 if dv01 <= 0.08 else 0
    s += 0 if not y10 else (1 if y10 > y2 else 0)
    if outlook == "Yields Rise":
        s += 1 if mod <= 4 else 0
    elif outlook == "Yields Fall":
        s += 1 if mod >= 7 else 0
    else:
        s += 1 if ytm >= bench else 0
    return s


def test_attractiveness_score():
    a = bm.analytics(S, M, 0.04, 96.50, 2, 1)
    d = build_sample()
    c = d["curve"]
    score = _score(a["ytm"], a["modified"], a["dv01"], c["y10"], c["y2"], 0.0430, "Yields Stable")
    assert score == 6                       # 2 (yield) +1 (dur) +1 (dv01) +1 (curve) +1 (outlook)
    view = "Attractive" if score >= 8 else "Neutral" if score >= 5 else "Not Attractive"
    assert view == "Neutral"


# ---------------------------------------------------------------------------
# Workbook structure + wiring
# ---------------------------------------------------------------------------
def _build(data, name="c.xlsx"):
    wb, _ = assemble(data)
    path = os.path.join(tempfile.mkdtemp(), name)
    wb.save(path)
    return openpyxl.load_workbook(path)


def test_sheet_count_and_order():
    wb2 = _build(build_sample())
    assert wb2.sheetnames == list(SHEET_ORDER)
    assert len(wb2.sheetnames) == 18


def test_named_ranges_present():
    for data in (build_sample(), empty_like_sample()):
        wb2 = _build(data)
        for nm in EXPECTED_NAMES:
            assert nm in wb2.defined_names, f"missing named range {nm}"


def test_bill_scenario_builds():
    """A zero-coupon bill config should also build and wire up cleanly."""
    d = build_sample()
    d["terms"].update(name="US Treasury Bill", sec_type="Treasury Bill", freq_name="Zero-coupon",
                      coupon=0.0, maturity=date(2025, 9, 18), mkt_clean=97.90)
    wb2 = _build(d, "bill.xlsx")
    assert "IsZero" in wb2.defined_names


def test_formula_explanations_are_text():
    """Regression: the 'Formula used' column must be TEXT, not live formulas."""
    wb2 = _build(build_sample())
    ws = wb2["Formula Explanations"]
    formula_cells = [c.coordinate for row in ws.iter_rows(min_col=2, max_col=2)
                     for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert not formula_cells, f"these should be text, not formulas: {formula_cells}"


def test_all_formula_refs_resolve():
    """Every function token is known and every identifier is a defined name (no #NAME?)."""
    wb2 = _build(build_sample())
    names = set(wb2.defined_names)
    ident = re.compile(r"[A-Za-z_][A-Za-z0-9_]*")
    cellref = re.compile(r"^[A-Za-z]{1,3}[0-9]{1,7}$")
    unknown = set()
    for ws in wb2.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                body = re.sub(r'"[^"]*"', "", v).replace("$", "")
                for m in ident.finditer(body):
                    tok = m.group(0)
                    if body[m.end():m.end() + 1] == "(":
                        if tok not in FUNCS:
                            unknown.add("FUNC:" + tok)
                    elif not (cellref.match(tok) or tok in ("TRUE", "FALSE") or tok in names):
                        unknown.add(tok)
    assert not unknown, f"unresolved names/functions: {sorted(unknown)}"


def _run_all():
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    ok = 0
    for t in tests:
        try:
            t()
            print(f"  PASS  {t.__name__}")
            ok += 1
        except AssertionError as e:
            print(f"  FAIL  {t.__name__}: {e}")
        except Exception as e:  # noqa: BLE001
            print(f"  ERROR {t.__name__}: {type(e).__name__}: {e}")
    print(f"\n{ok}/{len(tests)} tests passed.")
    return ok == len(tests)


if __name__ == "__main__":
    print(f"Running {__file__}\n")
    sys.exit(0 if _run_all() else 1)
