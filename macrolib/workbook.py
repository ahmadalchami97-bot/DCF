"""Builds the two-sheet Macro workbook: Impact Library + Global Stress Index."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from bond.style import apply as _styler
from dcf.utils import Refs, Sheet
from .data import BAND_INTERP, BAND_PORTFOLIO, STRESS_BANDS, build_sample

APP_NAME = "Macro Impact Library & Global Stress Index"

# arrow / word conditional-format colours (font only, keeps the sheet clean)
ARROWS = [("↑↑", "1F7A3D"), ("↑", "1F7A3D"), ("↓↓", "C00000"), ("↓", "C00000"),
          ("↔", "808080"), ("Mixed", "B7791F")]
FED = [("Hawkish", "C00000"), ("Dovish", "1F7A3D"), ("Neutral", "808080"), ("Mixed", "B7791F")]
CONF = [("High", "1F7A3D"), ("Medium", "B7791F"), ("Low", "808080")]
BAND_FILL = {"Calm": ("1F7A3D", "C6EFCE"), "Normal": ("1F7A3D", "E3F2E1"),
             "Elevated": ("9C6500", "FFEB9C"), "High Stress": ("C55A11", "FCE4D6"),
             "Crisis": ("9C0006", "FFC7CE")}
CALC_FILL = "F2F2F2"   # grey "looks protected" fill for formula cells


def _font_cf(ws, rng, pairs):
    for word, color in pairs:
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=[f'"{word}"'], font=Font(name="Calibri", size=10, bold=True, color=color)))


def _band_cf(ws, rng):
    for band, (fc, fl) in BAND_FILL.items():
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=[f'"{band}"'], font=Font(name="Calibri", size=11, bold=True, color=fc),
            fill=PatternFill(start_color=fl, end_color=fl, fill_type="solid")))


def _dropdown(sh, coord_range, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    sh.ws.add_data_validation(dv)
    dv.add(coord_range)


# ===========================================================================
# Sheet 1 - Macro Indicator Impact Library
# ===========================================================================
LIB_COLS = [
    ("Category", 14, "t"), ("Indicator", 22, "t"), ("Region", 12, "c"),
    ("What it measures", 26, "w"), ("Why it matters", 26, "w"),
    ("Higher-than-expected", 26, "w"), ("Lower-than-expected", 26, "w"),
    ("Inflation", 9, "a"), ("Growth", 9, "a"), ("Fed / CB", 12, "a"),
    ("USD / DXY", 9, "a"), ("UST yield", 9, "a"), ("UST bond px", 10, "a"),
    ("Gold", 8, "a"), ("Oil", 8, "a"), ("Equity / S&P", 11, "a"),
    ("EM", 8, "a"), ("Kuwait / GCC", 12, "a"),
    ("Usual market reaction", 30, "w"), ("Key exceptions / context", 30, "w"),
    ("Confidence", 11, "c"), ("Analyst explanation", 32, "w"),
]
LIB_KEYS = ["cat", "ind", "region", "meas", "why", "hi", "lo", "infl", "grow", "fed",
            "usd", "yld", "bond", "gold", "oil", "eq", "em", "gcc", "react", "exc", "conf", "note"]


def build_library(sh, inds):
    LAST = len(LIB_COLS)
    sh.hide_gridlines()
    sh.merge(1, 1, 1, LAST); sh.put(1, 1, "MACRO INDICATOR IMPACT LIBRARY", role="title"); sh.row_height(1, 26)
    sh.merge(2, 1, 2, LAST)
    sh.put(2, 1, "How major macro, market and geopolitical indicators usually affect markets - a professional cheat sheet",
           role="subtitle")
    sh.merge(3, 1, 3, LAST)
    sh.put(3, 1, "Impact columns = the USUAL first-order reaction to a HIGHER-than-expected print (data), a RISE (market "
                 "gauges) or ESCALATION (geopolitics).  ↑ up/supportive · ↓ down/negative · ↔ neutral · Mixed = depends "
                 "on context.  UST bond price moves opposite to yield.", role="note")
    sh.row_height(3, 26)

    hdr = 4
    for i, (name, _w, kind) in enumerate(LIB_COLS, 1):
        sh.put(hdr, i, name, role="colhdr" if kind in ("a", "c") else "colhdr_l", align="cw")
    sh.row_height(hdr, 30)

    r = hdr + 1
    for d in inds:
        for i, key in enumerate(LIB_KEYS, 1):
            kind = LIB_COLS[i - 1][2]
            val = d[key]
            if key in ("cat", "ind"):
                sh.put(r, i, val, role="label_b", align="lw" if key == "ind" else "l")
            elif kind == "w":
                sh.put(r, i, val, role="note_l", align="lw")
            elif kind in ("a", "c"):
                sh.put(r, i, val, role="label", align="c")
            else:
                sh.put(r, i, val, role="label", align="l")
        sh.row_height(r, 46)
        r += 1
    last = r - 1

    # conditional formatting: colour arrows, Fed words, confidence
    _font_cf(sh.ws, f"H5:R{last}", ARROWS)
    _font_cf(sh.ws, f"J5:J{last}", FED)
    _font_cf(sh.ws, f"U5:U{last}", CONF)
    _dropdown(sh, f"U5:U{last}", ["High", "Medium", "Low"])

    sh.ws.auto_filter.ref = f"A{hdr}:V{last}"
    sh.freeze("C5")
    for i, (_n, w, _k) in enumerate(LIB_COLS, 1):
        sh.col_width(i, w)
    return sh


# ===========================================================================
# Sheet 2 - Global Stress / Fear Index
# ===========================================================================
def _score_formula(driver, r):
    lvl = f"MEDIAN(0,100*(C{r}-F{r})/(G{r}-F{r}),100)"
    mov = f"MEDIAN(0,100*(E{r}-F{r})/(G{r}-F{r}),100)"
    neg = f"MEDIAN(0,100*(-E{r}-F{r})/(G{r}-F{r}),100)"
    ab = f"MEDIAN(0,100*(ABS(E{r})-F{r})/(G{r}-F{r}),100)"
    if driver == "level":
        return f'=IF(C{r}="","",IFERROR({lvl},0))'
    if driver == "move":
        return f'=IF(E{r}="","",IFERROR({mov},0))'
    if driver == "neg_move":
        return f'=IF(E{r}="","",IFERROR({neg},0))'
    if driver == "abs_move":
        return f'=IF(E{r}="","",IFERROR({ab},0))'
    return f'=IF(C{r}="","",MEDIAN(0,C{r},100))'   # manual


def _class_formula(cell):
    expr = '"Crisis"'
    for ub, lab in reversed(STRESS_BANDS):
        expr = f'IF({cell}<={ub},"{lab}",{expr})'
    return "=" + expr


def build_stress(sh, comps):
    LAST = 11
    sh.hide_gridlines()
    sh.merge(1, 1, 1, LAST); sh.put(1, 1, "GLOBAL STRESS / FEAR INDEX", role="title"); sh.row_height(1, 26)
    sh.merge(2, 1, 2, LAST)
    sh.put(2, 1, "A 0-100 composite of market fear and stress - are markets calm, normal, elevated, stressed or in crisis?",
           role="subtitle")
    sh.merge(3, 1, 3, LAST)
    sh.put(3, 1, "Yellow = you type (levels, changes, thresholds, weights). Grey = calculated. Score = Σ(component score × "
                 "weight) = SUMPRODUCT(scores, weights).", role="note")

    cfirst, clast = 10, 10 + len(comps) - 1
    srange, wrange, jrange, nrange = f"H{cfirst}:H{clast}", f"I{cfirst}:I{clast}", f"J{cfirst}:J{clast}", f"A{cfirst}:A{clast}"

    # --- headline ---
    sh.put(5, 1, "GLOBAL STRESS SCORE (0-100)", role="label_b"); sh.merge(5, 1, 5, 2)
    sh.put(5, 3, f"=ROUND(SUM({jrange}),0)", role="kpi"); sh.merge(5, 3, 5, 4)
    sh.put(5, 5, f"={_class_formula('C5')[1:]}", role="result", align="c"); sh.merge(5, 5, 5, 8)
    _band_cf(sh.ws, "E5:E5")
    sh.put(6, 1, "Main stress driver", role="label_b"); sh.merge(6, 1, 6, 2)
    sh.put(6, 3, f'=IFERROR(INDEX({nrange},MATCH(MAX({srange}),{srange},0)),"")', role="output_l"); sh.merge(6, 3, 6, 8)
    sh.put(7, 1, "Secondary stress driver", role="label_b"); sh.merge(7, 1, 7, 2)
    sh.put(7, 3, f'=IFERROR(INDEX({nrange},MATCH(LARGE({srange},2),{srange},0)),"")', role="output_l"); sh.merge(7, 3, 7, 8)

    # --- component table ---
    hdr = 9
    heads = ["Component", "Description", "Current level", "1-wk chg", "1-mo chg", "Normal thr.",
             "Stress thr.", "Score 0-100", "Weight", "Weighted", "Comment"]
    for i, h in enumerate(heads, 1):
        sh.put(hdr, i, h, role="colhdr" if 3 <= i <= 10 else "colhdr_l", align="cw")
    sh.row_height(hdr, 26)

    r = cfirst
    for c in comps:
        sh.put(r, 1, c["name"], role="label_b")
        sh.put(r, 2, c["desc"], role="note_l", align="lw")
        sh.put(r, 3, c["level"], role="input", fmt="0.0#")
        sh.put(r, 4, c["wk"], role="input", fmt="+0.0;-0.0;0.0")
        sh.put(r, 5, c["mo"], role="input", fmt="+0.0;-0.0;0.0")
        sh.put(r, 6, c["normal"], role="input", fmt="0.0#")
        sh.put(r, 7, c["stress"], role="input", fmt="0.0#")
        sh.put(r, 8, _score_formula(c["driver"], r), role="formula", fmt="0", fill=CALC_FILL)
        sh.put(r, 9, c["weight"], role="input", fmt="0%")
        sh.put(r, 10, f'=IF(H{r}="","",H{r}*I{r})', role="formula", fmt="0.0", fill=CALC_FILL)
        sh.put(r, 11, c["comment"], role="input_l")
        r += 1
    # totals / weight check
    sh.put(r, 1, "TOTAL", role="total_l"); sh.merge(r, 1, r, 7)
    sh.put(r, 8, f"=ROUND(SUM({jrange}),0)", role="total", fmt="0", fill=CALC_FILL)
    sh.put(r, 9, f"=SUM({wrange})", role="total", fmt="0%", fill=CALC_FILL)
    sh.put(r, 10, f"=SUM({jrange})", role="total", fmt="0.0", fill=CALC_FILL)
    sh.put(r, 11, f'=IF(ABS(SUM({wrange})-1)<0.001,"weights OK","WARNING: weights ≠ 100%")', role="status")
    _font_cf(sh.ws, f"K{r}:K{r}", [("weights OK", "1F7A3D")])
    tot_row = r

    # colour scale on the component scores
    sh.ws.conditional_formatting.add(srange, ColorScaleRule(
        start_type="num", start_value=0, start_color="C6EFCE",
        mid_type="num", mid_value=50, mid_color="FFEB9C",
        end_type="num", end_value=100, end_color="FFC7CE"))

    # --- interpretation (INDEX/MATCH into the legend below) ---
    lg_first = tot_row + 12
    lg_last = lg_first + 4
    bcol, icol, pcol = f"A{lg_first}:A{lg_last}", f"C{lg_first}:C{lg_last}", f"G{lg_first}:G{lg_last}"

    r = tot_row + 2
    sh.put(r, 1, "Market interpretation", role="label_b"); sh.merge(r, 1, r, 2)
    sh.put(r, 3, f'=IFERROR(INDEX({icol},MATCH(E5,{bcol},0)),"")', role="formula", align="lw", fill=CALC_FILL)
    sh.merge(r, 3, r, LAST); sh.row_height(r, 30)
    r += 1
    sh.put(r, 1, "Portfolio implication", role="label_b"); sh.merge(r, 1, r, 2)
    sh.put(r, 3, f'=IFERROR(INDEX({pcol},MATCH(E5,{bcol},0)),"")', role="formula", align="lw", fill=CALC_FILL)
    sh.merge(r, 3, r, LAST); sh.row_height(r, 30)
    r += 2
    sh.put(r, 1, "Auto summary", role="label_b"); sh.merge(r, 1, r, 2)
    sh.put(r, 3, '=IF(C5="","",'
                 '"Global market stress is "&LOWER(E5)&" ("&TEXT(C5,"0")&"/100). Main drivers: "&C6&" and "&C7&". "'
                 '&INDEX(' + icol + ',MATCH(E5,' + bcol + ',0))&" Portfolio: "&INDEX(' + pcol + ',MATCH(E5,' + bcol + ',0)))',
           role="output_l", align="lw", fill="E3F2E1")
    sh.merge(r, 3, r, LAST); sh.row_height(r, 56)
    r += 2
    sh.put(r, 1, "Analyst comment", role="label_b"); sh.merge(r, 1, r, 2)
    sh.put(r, 3, "", role="input_l"); sh.merge(r, 3, r, LAST); sh.row_height(r, 42)

    # --- classification legend (also the lookup source above) ---
    sh.put(lg_first - 2, 1, "Classification bands", role="section"); sh.merge(lg_first - 2, 1, lg_first - 2, LAST)
    hb = lg_first - 1
    for i, h in enumerate(["Band", "Range", "Market interpretation", "", "", "", "Portfolio implication"], 1):
        if h:
            sh.put(hb, i, h, role="colhdr_l")
    ranges = ["0-20", "21-40", "41-60", "61-80", "81-100"]
    bands = ["Calm", "Normal", "Elevated", "High Stress", "Crisis"]
    rr = lg_first
    for band, rng in zip(bands, ranges):
        sh.put(rr, 1, band, role="label_b")
        sh.put(rr, 2, rng, role="label", align="c")
        sh.put(rr, 3, BAND_INTERP[band], role="note_l", align="lw"); sh.merge(rr, 3, rr, 6)
        sh.put(rr, 7, BAND_PORTFOLIO[band], role="note_l", align="lw"); sh.merge(rr, 7, rr, LAST)
        sh.row_height(rr, 30)
        rr += 1
    _band_cf(sh.ws, f"A{lg_first}:A{lg_last}")

    widths = [22, 26, 13, 10, 10, 12, 12, 11, 9, 11, 30]
    for i, w in enumerate(widths, 1):
        sh.col_width(i, w)
    sh.freeze("A9")
    return sh


# ===========================================================================
def assemble(data: dict | None = None):
    data = data or build_sample()
    wb = Workbook()
    wb.remove(wb.active)
    refs = Refs()
    lib = Sheet(wb.create_sheet("Impact Library"), refs, styler=_styler)
    stress = Sheet(wb.create_sheet("Global Stress Index"), refs, styler=_styler)
    build_library(lib, data["indicators"])
    build_stress(stress, data["components"])
    for ws, color in ((lib.ws, "2E5984"), (stress.ws, "C00000")):
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.tabColor = color
    wb.active = 0
    wb.properties.title = APP_NAME
    return wb


def build_workbook(output_path: str, data: dict | None = None) -> str:
    assemble(data).save(output_path)
    return output_path
