"""
Styling layer.

Translates the design tokens in :mod:`dcf.config` into concrete openpyxl style
objects and exposes one workhorse function, :func:`apply`, that every sheet uses
to format a cell consistently. Centralising this guarantees a single, coherent
visual language across the whole workbook.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.protection import Protection

from .config import FONT_NAME, FONT_SIZE, FONT_SIZE_NOTE, FONT_SIZE_SMALL, Palette


# --- primitive builders -----------------------------------------------------
def _fill(hex_color: str | None) -> PatternFill | None:
    if not hex_color:
        return None
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def _font(color=Palette.TEXT, size=FONT_SIZE, bold=False, italic=False) -> Font:
    return Font(name=FONT_NAME, size=size, color=color, bold=bold, italic=italic)


_thin = Side(style="thin", color=Palette.GRID)
_medium = Side(style="medium", color=Palette.RULE)

BORDER_GRID = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
BORDER_BOX = Border(left=_medium, right=_medium, top=_medium, bottom=_medium)
BORDER_BOTTOM = Border(bottom=_thin)
BORDER_BOTTOM_RULE = Border(bottom=_medium)
BORDER_TOP_RULE = Border(top=_medium)
BORDER_NONE = Border()

ALIGN_L = Alignment(horizontal="left", vertical="center")
ALIGN_C = Alignment(horizontal="center", vertical="center")
ALIGN_R = Alignment(horizontal="right", vertical="center")
ALIGN_L_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_L_TOP_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
ALIGN_C_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)

LOCKED = Protection(locked=True)
UNLOCKED = Protection(locked=False)


# --- role definitions --------------------------------------------------------
# Each role maps to (font, fill, alignment, border). apply() can override any
# piece per call. Roles describe *meaning*, not raw appearance, so the visual
# language can be retuned in one place.
_ROLES = {
    "title": (_font(Palette.WHITE, 18, bold=True), _fill(Palette.NAVY_DARK), ALIGN_L, BORDER_NONE),
    "subtitle": (_font(Palette.WHITE, 11), _fill(Palette.NAVY_DARK), ALIGN_L, BORDER_NONE),
    "header": (_font(Palette.WHITE, 12, bold=True), _fill(Palette.NAVY), ALIGN_L, BORDER_NONE),
    "subheader": (_font(Palette.WHITE, 10, bold=True), _fill(Palette.SLATE), ALIGN_L, BORDER_NONE),
    "section": (_font(Palette.WHITE, 10, bold=True), _fill(Palette.STEEL), ALIGN_L, BORDER_NONE),
    "colhdr": (_font(Palette.WHITE, 10, bold=True), _fill(Palette.SLATE), ALIGN_C, BORDER_NONE),
    "colhdr_r": (_font(Palette.WHITE, 10, bold=True), _fill(Palette.SLATE), ALIGN_R, BORDER_NONE),
    "label": (_font(Palette.TEXT, FONT_SIZE), None, ALIGN_L, BORDER_NONE),
    "label_b": (_font(Palette.TEXT, FONT_SIZE, bold=True), None, ALIGN_L, BORDER_NONE),
    "sublabel": (_font(Palette.MUTED, FONT_SIZE_SMALL, italic=True), None, ALIGN_L, BORDER_NONE),
    "body": (_font(Palette.TEXT, FONT_SIZE), None, ALIGN_R, BORDER_NONE),
    "body_l": (_font(Palette.TEXT, FONT_SIZE), None, ALIGN_L, BORDER_NONE),
    "note": (_font(Palette.MUTED, FONT_SIZE_NOTE), None, ALIGN_L_WRAP, BORDER_NONE),
    "note_c": (_font(Palette.MUTED, FONT_SIZE_NOTE), None, ALIGN_C_WRAP, BORDER_NONE),
    # cell-type coding
    "input": (_font(Palette.INPUT_FONT, FONT_SIZE), _fill(Palette.INPUT_FILL), ALIGN_R, BORDER_GRID),
    "input_l": (_font(Palette.INPUT_FONT, FONT_SIZE), _fill(Palette.INPUT_FILL), ALIGN_L, BORDER_GRID),
    "calc": (_font(Palette.CALC_FONT, FONT_SIZE), None, ALIGN_R, BORDER_NONE),
    "link": (_font(Palette.LINK_FONT, FONT_SIZE), None, ALIGN_R, BORDER_NONE),
    "output": (_font(Palette.TEXT, FONT_SIZE, bold=True), _fill(Palette.OUTPUT_FILL), ALIGN_R, BORDER_GRID),
    "result": (_font(Palette.TEXT, 12, bold=True), _fill(Palette.RESULT_FILL), ALIGN_R, BORDER_BOX),
    # semantic status
    "good": (_font(Palette.GOOD, FONT_SIZE, bold=True), None, ALIGN_C, BORDER_NONE),
    "warn": (_font(Palette.WARN, FONT_SIZE, bold=True), None, ALIGN_C, BORDER_NONE),
    "bad": (_font(Palette.BAD, FONT_SIZE, bold=True), None, ALIGN_C, BORDER_NONE),
    "neutral": (_font(Palette.NEUTRAL, FONT_SIZE), None, ALIGN_C, BORDER_NONE),
    "panel": (_font(Palette.TEXT, FONT_SIZE), _fill(Palette.PANEL), ALIGN_L, BORDER_NONE),
}

_ALIGN_MAP = {
    "l": ALIGN_L, "c": ALIGN_C, "r": ALIGN_R,
    "lw": ALIGN_L_WRAP, "ltw": ALIGN_L_TOP_WRAP, "cw": ALIGN_C_WRAP,
}


def apply(
    cell,
    role: str = "body",
    *,
    fmt: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    align: str | None = None,
    fill: str | None = None,
    font_color: str | None = None,
    size: float | None = None,
    border=None,
    wrap: bool | None = None,
    locked: bool | None = None,
):
    """Apply a role's style to *cell*, with optional per-call overrides.

    Parameters mirror the common style axes. Any override that is ``None`` falls
    back to the role default, so call sites stay terse.
    """
    base_font, base_fill, base_align, base_border = _ROLES.get(role, _ROLES["body"])

    # Font (clone with overrides so we never mutate the cached object).
    f = base_font
    if any(v is not None for v in (bold, italic, font_color, size)):
        f = Font(
            name=FONT_NAME,
            size=size if size is not None else f.size,
            color=font_color if font_color is not None else f.color,
            bold=bold if bold is not None else f.bold,
            italic=italic if italic is not None else f.italic,
        )
    cell.font = f

    # Fill.
    if fill is not None:
        cell.fill = _fill(fill)
    elif base_fill is not None:
        cell.fill = base_fill

    # Alignment.
    if align is not None:
        cell.alignment = _ALIGN_MAP.get(align, base_align)
    elif wrap:
        cell.alignment = ALIGN_L_WRAP
    else:
        cell.alignment = base_align
    if wrap and align is None:
        cell.alignment = ALIGN_L_WRAP

    # Border.
    cell.border = border if border is not None else base_border

    # Number format.
    if fmt is not None:
        cell.number_format = fmt

    # Protection (workbook protection is opt-in per sheet).
    if locked is not None:
        cell.protection = LOCKED if locked else UNLOCKED

    return cell
