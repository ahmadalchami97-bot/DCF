"""
Shared context + drawing/visual helpers for the ratio workbook.

Reuses the workbook-agnostic engine from ``dcf.utils`` (the cross-sheet cell
registry and the styled ``Sheet`` wrapper) and adds this workbook's institutional
chrome: title/section bands, year headers (forecast years visually distinct),
hyperlink navigation, the legend, and the conditional-formatting visuals
(traffic lights, data bars, icon-set trend arrows, heat-map colour scales) plus
chart builders.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.chart import BarChart, LineChart, RadarChart, Reference
from openpyxl.formatting.rule import (ColorScaleRule, DataBarRule, FormulaRule,
                                      IconSetRule)
from openpyxl.styles import Font, PatternFill

from dcf.utils import Refs, Sheet  # reuse the proven engine
from .config import (APP_NAME, APP_VERSION, BAD_WORDS, FCST_YEARS, GOOD_WORDS,
                     HIST_YEARS, Palette, WARN_WORDS)

LABEL_COL = 1
FIRST_YEAR_COL = 2


@dataclass
class Context:
    data: dict
    refs: Refs
    sheets: dict = field(default_factory=dict)

    @property
    def periods(self):
        return self.data["periods"]

    @property
    def n_years(self):
        return len(self.periods)

    @property
    def hist(self):
        return self.data.get("hist_years", HIST_YEARS)

    @property
    def fcst(self):
        return self.data.get("fcst_years", FCST_YEARS)

    @property
    def last_hist_idx(self):
        return self.hist - 1

    @property
    def last_year_col(self):
        return FIRST_YEAR_COL + self.n_years - 1

    def year_col(self, i):
        return FIRST_YEAR_COL + i

    def sheet(self, name):
        return self.sheets[name]


# --- structural drawing -----------------------------------------------------
def title_block(sh: Sheet, title, subtitle, *, last_col, row=1):
    sh.merge(row, 1, row, last_col)
    sh.put(row, 1, title, role="title")
    sh.row_height(row, 30)
    sh.merge(row + 1, 1, row + 1, last_col)
    sh.put(row + 1, 1, subtitle, role="subtitle")
    sh.merge(row + 2, 1, row + 2, last_col)
    sh.put(row + 2, 1, f"{APP_NAME}  ·  v{APP_VERSION}", role="note")
    sh.hide_gridlines()
    return row + 3


def section(sh: Sheet, row, text, *, c1=1, c2=1, role="header"):
    if c2 > c1:
        sh.merge(row, c1, row, c2)
    sh.put(row, c1, text, role=role)
    sh.row_height(row, 18 if role != "header" else 20)
    return row + 1


def year_headers(sh: Sheet, ctx, row, *, label="Metric", extra=None):
    """Year header row; forecast years flagged with an 'F' suffix style."""
    sh.put(row, LABEL_COL, label, role="colhdr")
    for i, lab in enumerate(ctx.periods):
        is_fcst = i >= ctx.hist
        txt = f"{lab}{'  (F)' if is_fcst else ''}"
        c = sh.put(row, ctx.year_col(i), txt, role="colhdr_r")
        if is_fcst:
            c.font = Font(name="Calibri", size=9, bold=True, italic=True, color="CFE0EF")
    col = ctx.last_year_col
    if extra:
        for name in extra:
            col += 1
            sh.put(row, col, name, role="colhdr")
    sh.row_height(row, 22)
    return row + 1


def nav_bar(sh: Sheet, row, *, exclude=None, per_row=8, start_col=1):
    """Hyperlink 'buttons' to every sheet (no VBA needed)."""
    from .config import SHEET_ORDER
    exclude = exclude or set()
    targets = [s for s in SHEET_ORDER if s not in exclude]
    col = start_col
    for i, name in enumerate(targets):
        if i and i % per_row == 0:
            row += 1
            col = start_col
        c = sh.put(row, col, name, role="subheader", align="c")
        c.hyperlink = f"#'{name}'!A1"
        sh.col_width(col, 15)
        col += 1
    return row + 1


def legend(sh: Sheet, row, col=1):
    sh.put(row, col, "LEGEND", role="label_b")
    items = [
        ("input_l", "Input — editable (yellow)"),
        ("formula_l", "Formula / helper — protected (blue)"),
        ("output_l", "Output — read-only result (green)"),
        ("error", "Error / flag (red)"),
    ]
    r = row + 1
    for role, desc in items:
        sh.put(r, col, "  ", role=role)
        sh.put(r, col + 1, desc, role="note_l")
        sh.merge(r, col + 1, r, col + 4)
        r += 1
    sh.put(r, col, "Strong", role="status")
    sh.put(r, col + 1, "Moderate", role="status")
    sh.put(r, col + 2, "Weak", role="status")
    traffic_light(sh, f"{sh.coord(r, col)}:{sh.coord(r, col+2)}", sh.coord(r, col))
    sh.put(r, col + 3, "← traffic-light classification", role="note_l")
    sh.merge(r, col + 3, r, col + 6)
    return r + 1


def units_note(ctx):
    m = ctx.data["meta"]
    return f"All figures in {m.get('currency','USD')} {m.get('units','millions')} unless stated; (F) = forecast year."


# --- conditional-formatting visuals -----------------------------------------
def traffic_light(sh: Sheet, cell_range, anchor):
    """Colour classification text green/amber/red by sentiment keyword."""
    def rule(words):
        terms = ",".join(f'ISNUMBER(SEARCH("{w}",{anchor}))' for w in words)
        return [f"OR({terms})"]

    for words, font_c, fill_c in (
        (GOOD_WORDS, Palette.GOOD, Palette.GOOD_FILL),
        (WARN_WORDS, Palette.WARN, Palette.WARN_FILL),
        (BAD_WORDS, Palette.BAD, Palette.BAD_FILL),
    ):
        sh.ws.conditional_formatting.add(
            cell_range,
            FormulaRule(formula=rule(words), stopIfTrue=True,
                        font=Font(name="Calibri", size=10, bold=True, color=font_c),
                        fill=PatternFill(start_color=fill_c, end_color=fill_c, fill_type="solid")),
        )


def heat_map(sh: Sheet, cell_range, *, reverse=False):
    lo, hi = (Palette.BAD_FILL, Palette.GOOD_FILL)
    if reverse:
        lo, hi = hi, lo
    sh.ws.conditional_formatting.add(
        cell_range,
        ColorScaleRule(start_type="min", start_color=lo,
                       mid_type="percentile", mid_value=50, mid_color="FFF6C8",
                       end_type="max", end_color=hi))


def data_bars(sh: Sheet, cell_range, color=Palette.SUBHEADER):
    sh.ws.conditional_formatting.add(
        cell_range, DataBarRule(start_type="min", end_type="max", color=color))


def trend_icons(sh: Sheet, cell_range, *, reverse=False):
    rule = IconSetRule("3Arrows", "percent", [0, 33, 67], showValue=True, reverse=reverse)
    sh.ws.conditional_formatting.add(cell_range, rule)


# --- charts -----------------------------------------------------------------
def line_chart(sh: Sheet, anchor, title, *, rows, cat_row, first_col, last_col,
               width=20, height=8, style_id=2):
    ws = sh.ws
    chart = LineChart()
    chart.title = title
    chart.style = style_id
    chart.height = height
    chart.width = width
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    data = Reference(ws, min_col=LABEL_COL, max_col=last_col,
                     min_row=rows[0], max_row=rows[-1])
    chart.add_data(data, from_rows=True, titles_from_data=True)
    cats = Reference(ws, min_col=first_col, max_col=last_col, min_row=cat_row, max_row=cat_row)
    chart.set_categories(cats)
    for s in chart.series:
        s.smooth = False
    ws.add_chart(chart, anchor)


def bar_chart(sh: Sheet, anchor, title, *, rows, cat_row, first_col, last_col,
              width=20, height=8, style_id=10):
    ws = sh.ws
    chart = BarChart()
    chart.type = "col"
    chart.title = title
    chart.style = style_id
    chart.height = height
    chart.width = width
    data = Reference(ws, min_col=LABEL_COL, max_col=last_col, min_row=rows[0], max_row=rows[-1])
    chart.add_data(data, from_rows=True, titles_from_data=True)
    cats = Reference(ws, min_col=first_col, max_col=last_col, min_row=cat_row, max_row=cat_row)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)


def radar_chart(sh: Sheet, anchor, title, *, rows, cat_row, first_col, last_col,
                width=12, height=10):
    ws = sh.ws
    chart = RadarChart()
    chart.type = "filled"
    chart.title = title
    chart.height = height
    chart.width = width
    data = Reference(ws, min_col=LABEL_COL, max_col=last_col, min_row=rows[0], max_row=rows[-1])
    chart.add_data(data, from_rows=True, titles_from_data=True)
    cats = Reference(ws, min_col=first_col, max_col=last_col, min_row=cat_row, max_row=cat_row)
    chart.set_categories(cats)
    ws.add_chart(chart, anchor)
