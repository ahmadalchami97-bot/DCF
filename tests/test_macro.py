"""
Test suite for the Macro Market Intelligence System (core engine).

The scoring is transparent arithmetic, so we re-derive it in Python from the same
seed data (mirroring the Excel formulas) and assert the model behaves as designed:
the surprise/polarity logic (a lower-unemployment beat reads hawkish), and the Fed
composite lands in the expected band. We also check the workbook builds, wires up
its named ranges, keeps Fed weights summing to 1.0, and that every formula name /
function resolves (no #NAME?).

Run: ``python tests/test_macro.py``  (or ``pytest -q tests/test_macro.py``).
"""

from __future__ import annotations

import os
import re
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402

from macro.config import FED_SUBSCORES, SHEET_ORDER  # noqa: E402
from macro.data import (FED_SPEECHES, FED_STATE, INDICATORS, SAMPLE_RELEASES,  # noqa: E402
                        build_sample, empty_like_sample)
from macro.workbook import assemble  # noqa: E402

LIB = {d["name"]: d for d in INDICATORS}
FUNCS = set((
    "IF IFERROR IFNA SUM SUMPRODUCT MIN MAX ABS SIGN N ISNUMBER NOT OR AND COUNTIF SEARCH CHOOSE TEXT ROUND "
    "INDEX MATCH AVERAGEIF AVERAGE MEDIAN TODAY LEFT HYPERLINK"
).split())
EXPECTED_NAMES = (
    "Lib_Name", "Lib_Cat", "Lib_Interp_Higher", "Lib_Interp_Lower", "Lib_PolInfl", "Lib_PolGrow",
    "Lib_Importance", "Lib_Scale", "SUR_CAT", "SUR_INFLIMP", "SUR_GROWIMP",
    "Fed_Weights", "Fed_SubScores", "Score_Fed", "Fed_Bias_Label",
)


# ---------------------------------------------------------------------------
# The scoring model, re-derived in Python (mirrors the Excel formulas)
# ---------------------------------------------------------------------------
def _clamp(x):
    return max(-2.0, min(2.0, x))


def _impulses():
    rows = []
    for (_dt, _c, ind, act, fc, _prev, *_rest) in SAMPLE_RELEASES:
        d = LIB[ind]
        z = (act - fc) / d["scale"]
        rows.append(dict(ind=ind, cat=d["cat"], z=z,
                         infl=z * d["pinf"] * d["imp"] / 5, grow=z * d["pgro"] * d["imp"] / 5))
    return rows


def _ladder(x):
    return 2 if x >= 15 else 1 if x >= 5 else 0 if x > -5 else -1 if x > -15 else -2


def _fed_composite():
    rows = _impulses()

    def avg(cat, key):
        vals = [r[key] for r in rows if r["cat"] == cat]
        return sum(vals) / len(vals) if vals else 0.0

    tones = [s[2] for s in FED_SPEECHES]
    subs = [_clamp(avg("Inflation", "infl")), _clamp(avg("Labor", "infl")), _clamp(avg("Growth", "grow")),
            _clamp(sum(tones) / len(tones) if tones else 0.0),
            _ladder(FED_STATE["implied_bps"]), _ladder(FED_STATE["curve_2y_bps"]),
            _ladder(FED_STATE["real_10y_bps"])]
    weights = [w for _lbl, w in FED_SUBSCORES]
    return subs, sum(s * w for s, w in zip(subs, weights))


def test_surprise_z_and_polarity():
    rows = {r["ind"]: r for r in _impulses()}
    # a hot CPI beat: z = +1, inflationary impulse positive
    assert abs(rows["CPI y/y"]["z"] - 1.0) < 1e-9
    assert rows["CPI y/y"]["infl"] > 0
    # THE polarity test: unemployment *fell* vs forecast (a beat) -> hawkish + stronger growth
    assert rows["Unemployment rate"]["z"] < 0            # raw surprise is negative
    assert rows["Unemployment rate"]["infl"] > 0         # but the impulse is hawkish
    assert rows["Unemployment rate"]["grow"] > 0         # and growth-positive
    # fewer jobless claims (a beat, claims fell) also reads hawkish/strong
    assert rows["Initial jobless claims"]["z"] < 0 and rows["Initial jobless claims"]["infl"] > 0


def test_fed_weights_sum_to_one():
    assert abs(sum(w for _l, w in FED_SUBSCORES) - 1.0) < 1e-9


def test_fed_composite_is_mildly_hawkish():
    subs, comp = _fed_composite()
    assert len(subs) == 7
    assert 0.33 <= comp < 1.0, f"expected mildly hawkish, got {comp:.3f}"
    # inflation and labor sub-scores should both be hawkish on this tape
    assert subs[0] > 0 and subs[1] > 0


def test_growth_subscore_is_mixed():
    subs, _ = _fed_composite()
    # firm ISM services / retail offset by soft ISM mfg, confidence, housing -> near flat
    assert -0.5 < subs[2] < 0.5


# ---------------------------------------------------------------------------
# Workbook structure + wiring
# ---------------------------------------------------------------------------
def _build(data, name="m.xlsx"):
    wb, _ = assemble(data)
    path = os.path.join(tempfile.mkdtemp(), name)
    wb.save(path)
    return openpyxl.load_workbook(path)


def test_sheet_count_and_order():
    wb2 = _build(build_sample())
    assert wb2.sheetnames == list(SHEET_ORDER)
    assert len(wb2.sheetnames) == 6


def test_named_ranges_present_both_variants():
    for data in (build_sample(), empty_like_sample()):
        wb2 = _build(data)
        for nm in EXPECTED_NAMES:
            assert nm in wb2.defined_names, f"missing named range {nm}"


def test_subscores_and_weights_same_size():
    """SUMPRODUCT(Fed_SubScores, Fed_Weights) needs equal-length ranges."""
    wb2 = _build(build_sample())

    def span(name):
        ref = wb2.defined_names[name].value          # e.g. "'Fed Tracker'!$C$40:$C$46"
        rng = ref.split("!")[1].replace("$", "")
        a, b = rng.split(":")
        return int(re.sub(r"[A-Za-z]", "", b)) - int(re.sub(r"[A-Za-z]", "", a)) + 1
    assert span("Fed_SubScores") == 7
    assert span("Fed_Weights") == 7


def test_all_formula_refs_resolve():
    wb2 = _build(build_sample())
    names = set(wb2.defined_names)
    ident = re.compile(r"[A-Za-z_][A-Za-z0-9_]*")
    cellref = re.compile(r"^[A-Za-z]{1,3}[0-9]{1,7}$")
    unknown = set()
    for ws in wb2.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                body = re.sub(r'"[^"]*"', "", v)
                body = re.sub(r"'[^']*'", "", body).replace("$", "")   # drop quoted sheet names
                for m in ident.finditer(body):
                    tok = m.group(0)
                    if body[m.end():m.end() + 1] == "(":
                        if tok not in FUNCS:
                            unknown.add("FUNC:" + tok)
                    elif not (cellref.match(tok) or tok in ("TRUE", "FALSE") or tok in names):
                        unknown.add(tok)
    assert not unknown, f"unresolved names/functions: {sorted(unknown)}"


def test_blank_has_no_sample_rows():
    """The blank template clears releases/speeches but keeps the reference library."""
    wb2 = _build(empty_like_sample(), "blank.xlsx")
    ws = wb2["Data Tracker"]
    # first data row's indicator cell (D9) should be empty in the blank build
    assert ws["D9"].value in (None, "")
    assert "Lib_Name" in wb2.defined_names   # library still present


def _run_all():
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    ok = 0
    for t in tests:
        try:
            t()
            print(f"  PASS  {t.__name__}")
            ok += 1
        except AssertionError as e:
            print(f"  FAIL  {t.__name__}: {e}")
        except Exception as e:  # noqa: BLE001
            print(f"  ERROR {t.__name__}: {type(e).__name__}: {e}")
    print(f"\n{ok}/{len(tests)} tests passed.")
    return ok == len(tests)


if __name__ == "__main__":
    print(f"Running {__file__}\n")
    sys.exit(0 if _run_all() else 1)
