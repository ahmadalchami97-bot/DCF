"""
Test suite for the two-sheet Macro workbook (macrolib).

Checks the library is complete and well-formed, that the Global Stress Index
scoring (re-derived in Python from the same seed) lands where expected, that the
workbook has exactly two sheets, and that every formula name/function resolves.

Run: ``python tests/test_macrolib.py``  (or ``pytest -q tests/test_macrolib.py``).
"""

from __future__ import annotations

import os
import re
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402

from macrolib.data import COMPONENTS, INDICATORS, STRESS_BANDS, build_sample, empty_like_sample  # noqa: E402
from macrolib.workbook import LIB_KEYS, assemble  # noqa: E402

FUNCS = set("IF IFERROR ABS SUM SUMPRODUCT MEDIAN INDEX MATCH MAX LARGE ROUND TEXT LOWER N".split())


def _med(a, x, b):
    return sorted([a, x, b])[1]


def _score(c):
    lo, hi, d = c["normal"], c["stress"], c["driver"]
    if d == "manual":
        return _med(0, c["level"], 100)
    m = {"level": c["level"], "move": c["mo"], "neg_move": -c["mo"], "abs_move": abs(c["mo"])}[d]
    return _med(0, 100 * (m - lo) / (hi - lo), 100)


# ---------------------------------------------------------------------------
# Library
# ---------------------------------------------------------------------------
def test_library_is_large_and_complete():
    assert len(INDICATORS) >= 60
    for d in INDICATORS:
        for k in LIB_KEYS:
            assert k in d and d[k] != "" or k in ("react", "exc", "note"), f"{d.get('ind')} missing {k}"


def test_key_indicators_present():
    names = {d["ind"] for d in INDICATORS}
    for must in ("CPI", "Core PCE", "Nonfarm Payrolls", "FOMC Decision", "DXY (US Dollar Index)",
                 "US 10Y Yield", "VIX", "MOVE Index", "Gold", "Brent Oil", "Strait of Hormuz Risk",
                 "OPEC+ Decisions", "China GDP", "ECB Policy", "BOJ Policy"):
        assert must in names, f"missing indicator {must}"


def test_bond_price_inverts_yield():
    inv = {"↑": "↓", "↓": "↑", "↑↑": "↓↓", "↓↓": "↑↑"}
    for d in INDICATORS:
        if d["yld"] in inv:
            assert d["bond"] == inv[d["yld"]], f"{d['ind']}: bond should invert yield"


# ---------------------------------------------------------------------------
# Stress index
# ---------------------------------------------------------------------------
def test_weights_sum_to_one():
    assert abs(sum(c["weight"] for c in COMPONENTS) - 1.0) < 1e-9


def test_nine_components():
    assert len(COMPONENTS) == 9


def test_stress_score_is_elevated():
    total = round(sum(_score(c) * c["weight"] for c in COMPONENTS))
    band = next((lab for ub, lab in STRESS_BANDS if total <= ub), "Crisis")
    assert 41 <= total <= 60 and band == "Elevated", f"got {total} / {band}"


def test_main_driver_is_usd_then_oil():
    ranked = sorted(COMPONENTS, key=_score, reverse=True)
    assert "DXY" in ranked[0]["name"]
    assert "Brent" in ranked[1]["name"] or "Geopolitical" in ranked[1]["name"]


def test_scores_bounded_0_100():
    for c in COMPONENTS:
        assert 0 <= _score(c) <= 100


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------
def _build(data, name="ml.xlsx"):
    wb = assemble(data)
    path = os.path.join(tempfile.mkdtemp(), name)
    wb.save(path)
    return openpyxl.load_workbook(path)


def test_exactly_two_sheets():
    for data in (build_sample(), empty_like_sample()):
        wb2 = _build(data)
        assert wb2.sheetnames == ["Impact Library", "Global Stress Index"]


def test_all_formula_refs_resolve():
    wb2 = _build(build_sample())
    names = set(wb2.defined_names)
    ident = re.compile(r"[A-Za-z_][A-Za-z0-9_]*")
    cellref = re.compile(r"^[A-Za-z]{1,3}[0-9]{1,7}$")
    unknown = set()
    for ws in wb2.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                body = re.sub(r'"[^"]*"', "", v)
                body = re.sub(r"'[^']*'", "", body).replace("$", "")
                for m in ident.finditer(body):
                    tok = m.group(0)
                    if body[m.end():m.end() + 1] == "(":
                        if tok not in FUNCS:
                            unknown.add("FUNC:" + tok)
                    elif not (cellref.match(tok) or tok in ("TRUE", "FALSE") or tok in names):
                        unknown.add(tok)
    assert not unknown, f"unresolved names/functions: {sorted(unknown)}"


def _run_all():
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    ok = 0
    for t in tests:
        try:
            t()
            print(f"  PASS  {t.__name__}")
            ok += 1
        except AssertionError as e:
            print(f"  FAIL  {t.__name__}: {e}")
        except Exception as e:  # noqa: BLE001
            print(f"  ERROR {t.__name__}: {type(e).__name__}: {e}")
    print(f"\n{ok}/{len(tests)} tests passed.")
    return ok == len(tests)


if __name__ == "__main__":
    print(f"Running {__file__}\n")
    sys.exit(0 if _run_all() else 1)
